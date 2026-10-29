"""The hk helper library, run inside the sandbox image (`python -m pytest /opt/hk/tests`)."""

from __future__ import annotations

import datetime as dt
import io
import json
import re
import time
import zipfile

import hk
import numpy as np
import openpyxl
import pandas as pd
import pytest
from conftest import image_pdf, merge_pdfs, text_pdf
from hk._inspect import inspect


def by_type(records, kind):
    return [r for r in records if r["type"] == kind]


# ------------------------------------------------------------------ reading


def test_read_csv_bom_semicolon_decimal_comma(job):
    data = "﻿Date;Amount\r\n2026-01-05;1.234,5\r\n2026-01-06;12,00\r\n2026-01-07;3\r\n"
    fid = job.add("statement.csv", "csv", data.encode("utf-8"))
    df = hk.read(fid)
    assert list(df.columns) == ["Date", "Amount"]
    assert df["Amount"].dtype == float
    assert df["Amount"].tolist() == [1234.5, 12.0, 3.0]
    assert pd.api.types.is_datetime64_any_dtype(df["Date"])
    assert df.attrs["hk_first_data_row"] == 2
    assert hk.read("statement.csv").equals(df)  # by name as well as by id


def test_read_csv_cp1252_and_tabs(job):
    data = "Name\tPrice\nCafé\t2.50\nCrème\t3.00\n".encode("cp1252")
    df = hk.read(job.add("menu.txt", "txt", data))
    assert list(df.columns) == ["Name", "Price"]
    assert df["Name"].tolist() == ["Café", "Crème"]
    assert df["Price"].tolist() == [2.5, 3.0]


def test_read_finds_header_not_on_line_1(job):
    data = (
        "Statement for account ending 1234\n"
        "Period: January 2026\n"
        "\n"
        "Date,Description,Amount\n"
        "2026-01-05,Coffee,-3.20\n"
        "2026-01-06,Salary,2500.00\n"
    )
    df = hk.read(job.add("bank.csv", "csv", data.encode()))
    assert list(df.columns) == ["Date", "Description", "Amount"]
    assert len(df) == 2
    assert df["Amount"].tolist() == [-3.2, 2500.0]
    assert df.attrs["hk_header_row"] == 4
    assert df.attrs["hk_first_data_row"] == 5


def test_read_xlsx_values_not_formulas(job, tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Amount", "Double"])
    ws.append([42, "=A2*2"])
    raw = io.BytesIO()
    wb.save(raw)
    # Excel stores the formula's last result next to it; openpyxl doesn't, so add it.
    src = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    patched = io.BytesIO()
    with zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = body.decode("utf-8")
                xml, n = re.subn(r"<f>A2\*2</f>(<v\s*/>|<v></v>)?", "<f>A2*2</f><v>84</v>", xml)
                assert n == 1
                body = xml.encode("utf-8")
            dst.writestr(item, body)
    df = hk.read(job.add("book.xlsx", "xlsx", patched.getvalue()))
    assert list(df.columns) == ["Amount", "Double"]
    assert df["Double"].tolist() == [84]
    assert not any(isinstance(v, str) and "A2" in v for v in df.to_numpy().ravel())


def test_read_large_sheet_streams(job, tmp_path):
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Data")
    ws.append(["Id", "Amount", "Category"])
    for i in range(200_000):
        ws.append([i, i * 0.5, "abc"[i % 3]])
    path = tmp_path / "big.xlsx"
    wb.save(path)
    fid = job.add("big.xlsx", "xlsx", path.read_bytes())
    started = time.monotonic()
    df = hk.read(fid)
    assert time.monotonic() - started < 60
    assert len(df) == 200_000
    assert df["Amount"].iloc[-1] == 199_999 * 0.5


def test_read_json_flattens(job):
    data = json.dumps({"items": [{"id": 1, "price": {"gbp": 2.5}}, {"id": 2, "price": {"gbp": 4}}]})
    df = hk.read(job.add("items.json", "json", data.encode()))
    assert list(df.columns) == ["id", "price.gbp"]
    assert len(df) == 2


def test_read_unknown_file_lists_the_files(job):
    job.add("a.csv", "csv", b"x,y\n1,2\n")
    with pytest.raises(ValueError, match="a.csv"):
        hk.read("nope")


# ------------------------------------------------------------------ reporting


def test_result_records_numbers_dates_and_nan(job):
    hk.result("count", np.int64(5))
    hk.result("total", 2.5, unit="GBP", note="after refunds")
    hk.result("first day", dt.date(2026, 1, 5))
    hk.result("last payment", pd.Timestamp("2026-01-31 10:00"))
    hk.result("average", float("nan"))
    hk.result("growth", np.float64("inf"))
    hk.document("closing balance", 1234.56, page=2, unit="GBP")
    values = {r["name"]: r for r in by_type(job.results(), "value")}
    assert values["count"]["value"] == 5 and values["count"]["is_date"] is False
    assert values["count"]["kind"] == "computed" and values["count"]["page"] is None
    assert values["total"]["unit"] == "GBP" and values["total"]["note"] == "after refunds"
    assert values["first day"] == {
        "type": "value",
        "name": "first day",
        "value": "2026-01-05",
        "unit": None,
        "note": None,
        "is_date": True,
        "kind": "computed",
        "page": None,
    }
    assert values["last payment"]["value"].startswith("2026-01-31T10:00")
    assert values["last payment"]["is_date"] is True
    assert values["average"]["value"] is None and values["average"]["note"] == "not a number"
    assert values["growth"]["value"] is None and values["growth"]["note"] == "not a number"
    assert values["closing balance"]["kind"] == "document"
    assert values["closing balance"]["page"] == 2
    with pytest.raises(TypeError):
        hk.result("rows", [1, 2, 3])


def test_table_caps_rows_and_keeps_file_row_numbers(job):
    data = "n,label\n" + "".join(f"{i},row {i}\n" for i in range(800))
    df = hk.read(job.add("rows.csv", "csv", data.encode()))
    hk.table("All rows", df, note="every row")
    hk.table("Late rows", df[df["n"] >= 700], columns=["n"])
    first, late = by_type(job.results(), "table")
    assert first["total_rows"] == 800
    assert len(first["rows"]) == 500
    assert first["columns"] == ["n", "label"]
    assert first["row_numbers"][0] == 2
    assert first["row_numbers"][-1] == 501
    assert first["rows"][0] == [0, "row 0"]
    assert first["note"] == "every row"
    assert late["total_rows"] == 100
    assert late["columns"] == ["n"]
    assert late["row_numbers"][0] == 702
    assert late["rows"][0] == [700]


def _monthly():
    return pd.DataFrame(
        {
            "Month": pd.to_datetime(["2026-01-01", "2026-02-01", "2026-03-01"] * 2),
            "Amount": [10.0, 12.5, 9.0, 4.0, 5.0, np.nan],
            "Category": ["Food"] * 3 + ["Travel"] * 3,
        }
    )


def test_bar_line_scatter_specs_have_zoom_legend_and_highlight(job):
    df = _monthly()
    highlight = [
        {"x": "2026-02-01", "label": "Moved house"},
        {"from": "2026-01-01", "to": "2026-01-15", "label": "Holiday"},
    ]
    hk.line(df, "Month", "Amount", color="Category", title="By month", highlight=highlight)
    hk.scatter(df, "Month", "Amount", title="Points")
    hk.bar(df, "Category", "Amount", color="Category", title="Totals")
    line, scatter, bar = by_type(job.results(), "chart")

    spec = line["spec"]
    assert line["title"] == "By month"
    assert spec["$schema"] == "https://vega.github.io/schema/vega-lite/v6.json"
    assert len(spec["data"]["values"]) == 6
    assert spec["data"]["values"][0] == {
        "Month": "2026-01-01T00:00:00",
        "Amount": 10.0,
        "Category": "Food",
    }
    assert spec["data"]["values"][5]["Amount"] is None
    main = spec["layer"][0]
    assert main["mark"]["type"] == "line"
    assert main["encoding"]["x"] == {"field": "Month", "type": "temporal", "title": "Month"}
    params = {p["name"]: p for p in main["params"]}
    assert params["zoom"] == {"name": "zoom", "select": "interval", "bind": "scales"}
    assert params["legend"]["bind"] == "legend"
    assert params["legend"]["select"] == {"type": "point", "fields": ["Category"]}
    assert main["encoding"]["opacity"]["condition"]["param"] == "legend"
    marks = [layer["mark"]["type"] for layer in spec["layer"][1:]]
    assert "rule" in marks and "rect" in marks
    rect = next(layer for layer in spec["layer"] if layer["mark"]["type"] == "rect")
    assert rect["data"]["values"][0]["hk_label"] == "Holiday"
    json.dumps(spec, allow_nan=False)

    scatter_main = scatter["spec"]["layer"][0]
    assert scatter_main["mark"]["type"] == "point"
    assert [p["name"] for p in scatter_main["params"]] == ["zoom"]
    assert "color" not in scatter_main["encoding"]

    bar_main = bar["spec"]["layer"][0]
    assert bar_main["mark"]["type"] == "bar"
    assert [p["name"] for p in bar_main["params"]] == ["legend"]

    with pytest.raises(ValueError):
        hk.line(df, "Month", "Amount", title="Bad", highlight=[{"label": "nowhere"}])


def test_chart_rejects_over_5000_rows(job):
    df = pd.DataFrame({"x": range(5001), "y": range(5001)})
    with pytest.raises(ValueError, match="5,000"):
        hk.scatter(df, "x", "y", title="Too many")
    with pytest.raises(ValueError, match="5,000"):
        hk.chart(
            {"data": {"values": [{"x": i} for i in range(5001)]}, "mark": "point"}, title="Too many"
        )
    hk.scatter(df.head(5000), "x", "y", title="Just right")
    assert len(by_type(job.results(), "chart")) == 1


def test_file_writes_csv_and_xlsx(job):
    df = pd.DataFrame({"Name": ["Café", "Tea"], "Price": [2.5, 1.0]})
    hk.file("prices.csv", df)
    hk.file("prices.xlsx", df)
    hk.file("notes.txt", "hello")
    hk.file("raw.bin", b"\x00\x01")
    csv_bytes = (job.out / "files" / "prices.csv").read_bytes()
    assert csv_bytes.startswith(b"\xef\xbb\xbfName,Price")
    assert "Café".encode() in csv_bytes
    wb = openpyxl.load_workbook(job.out / "files" / "prices.xlsx")
    assert [c.value for c in wb.active[1]] == ["Name", "Price"]
    assert (job.out / "files" / "notes.txt").read_text(encoding="utf-8") == "hello"
    assert (job.out / "files" / "raw.bin").read_bytes() == b"\x00\x01"
    members = [r["member"] for r in by_type(job.results(), "file")]
    assert members == ["files/prices.csv", "files/prices.xlsx", "files/notes.txt", "files/raw.bin"]
    for bad in ("../escape.csv", ".hidden", "a/b.csv", "", "x" * 101, "tab\tname.csv"):
        with pytest.raises(ValueError):
            hk.file(bad, "x")
    hk.cleaned("Dropped the totals row.")
    assert by_type(job.results(), "cleaning") == [
        {"type": "cleaning", "note": "Dropped the totals row."}
    ]


# ------------------------------------------------------------------ profile


def test_profile_reports_glance_and_problems(job):
    data = (
        "Date,Description,Amount\n"
        "2026-01-05,Coffee,-3.20\n"
        "05/01/2026,Rent,-950.00\n"
        "2026-01-07,Salary,2500.00\n"
        "2026-01-08,Refund,-12.00\n"
        "2026-01-09,Book,15.00\n"
        "2026-01-09,Book,15.00\n"
        "2026-01-10,Gift,40.00\n"
        "2026-01-10,Gift,40.00\n"
        "2026-01-11,,7.50\n"
    )
    hk.profile(job.add("bank.csv", "csv", data.encode()))
    records = job.results()
    assert {"type": "section", "name": "glance"} in records
    values = {r["name"]: r["value"] for r in by_type(records, "value")}
    assert values["rows"] == 9
    assert values["columns"] == 3
    assert values["missing in Description"] == 1
    notes = " | ".join(r["note"] for r in by_type(records, "cleaning"))
    assert "mixes date formats" in notes
    assert "3 negative values" in notes
    assert "2 rows are exact copies" in notes
    tables = {t["name"]: t for t in by_type(records, "table")}
    assert tables["Negative values in Amount"]["total_rows"] == 3
    assert tables["Negative values in Amount"]["row_numbers"] == [2, 3, 5]
    assert tables["Duplicate rows"]["total_rows"] == 2
    assert tables["Duplicate rows"]["row_numbers"] == [7, 9]
    assert tables["Dates not in the usual format in Date"]["row_numbers"] == [3]


def test_profile_notes_header_and_numbers_stored_as_text(job):
    data = (
        "Exported from online banking\n"
        "Date,Payee,Amount\n"
        '2026-01-05,Shop,"1,234.50"\n'
        '2026-01-06,Cafe,"12.00"\n'
        '2026-01-07,Rent,"2,000.00"\n'
    )
    hk.profile(job.add("export.csv", "csv", data.encode()))
    notes = [r["note"] for r in by_type(job.results(), "cleaning")]
    assert any("line 2, not line 1" in n for n in notes)
    assert any("numbers stored as text" in n for n in notes)
    assert any("thousands separators" in n for n in notes)


# ------------------------------------------------------------------ PDFs and inspection

STATEMENT = ["Example Bank", "Opening balance 100.00", "Closing balance 1234.56"]


def test_inspect_owner_password_only_pdf_opens(job):
    pdf = merge_pdfs(text_pdf(STATEMENT), user_password="", owner_password="owner-only")
    fid = job.add("statement.pdf", "pdf", pdf)
    path = job.manifest["files"][0]["path"]
    report = inspect(path, "pdf")
    assert report["ok"] is True
    assert report["problem"] is None
    assert report["pages"] == 1
    assert report["scanned_pages"] == []
    pages = hk.pdf_text(fid)
    assert pages[0]["page"] == 1
    assert "Closing balance 1234.56" in pages[0]["text"]


def test_inspect_user_password_pdf_refused(job):
    pdf = merge_pdfs(text_pdf(STATEMENT), user_password="open-sesame", owner_password="x")
    job.add("locked.pdf", "pdf", pdf)
    report = inspect(job.manifest["files"][0]["path"], "pdf")
    assert report["ok"] is False
    assert report["problem"] == "password"


def test_inspect_scanned_pdf_flags_pages(job):
    pdf = merge_pdfs(text_pdf(STATEMENT), image_pdf(), text_pdf(["Page three"]))
    job.add("mixed.pdf", "pdf", pdf)
    report = inspect(job.manifest["files"][0]["path"], "pdf")
    assert report["ok"] is True
    assert report["pages"] == 3
    assert report["scanned_pages"] == [2]


def test_inspect_rc4_xls_is_password(job, monkeypatch):
    # A legacy .xls encrypted with RC4 passes the attach check (FILEPASS record, no
    # EncryptedPackage stream); xlrd refuses it with an error that mentions encryption.
    import xlrd

    def encrypted(*args, **kwargs):
        raise xlrd.XLRDError("Workbook is encrypted")

    monkeypatch.setattr(xlrd, "open_workbook", encrypted)
    job.add("old.xls", "xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504)
    report = inspect(job.manifest["files"][0]["path"], "xls")
    assert report["ok"] is False
    assert report["problem"] == "password"


def test_inspect_damaged_xls_is_unreadable(job):
    job.add("bad.xls", "xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504)
    report = inspect(job.manifest["files"][0]["path"], "xls")
    assert report == {**report, "ok": False, "problem": "unreadable"}


def test_inspect_csv_reports_encoding_delimiter_and_columns(job):
    data = "﻿Date;Amount\n2026-01-05;1,5\n2026-01-06;2,5\n"
    job.add("s.csv", "csv", data.encode("utf-8"))
    report = inspect(job.manifest["files"][0]["path"], "csv")
    assert report["ok"] is True
    assert report["encoding"] == "utf-8-sig"
    assert report["delimiter"] == ";"
    assert report["rows"] == 2
    assert report["columns"] == ["Date", "Amount"]


def test_inspect_xlsx_lists_sheets(job, tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Spending"
    wb.active.append(["Report"])
    wb.active.append(["Date", "Amount"])
    for i in range(5):
        wb.active.append([dt.date(2026, 1, i + 1), i])
    wb.create_sheet("Notes").append(["hello"])
    path = tmp_path / "book.xlsx"
    wb.save(path)
    report = inspect(path, "xlsx")
    assert report["ok"] is True
    assert report["sheets"] == [
        {"name": "Spending", "rows": 7, "columns": 2},
        {"name": "Notes", "rows": 1, "columns": 1},
    ]
    assert report["columns"] == ["Date", "Amount"]
    assert report["rows"] == 5


def test_pdf_helpers_need_a_pdf(job):
    fid = job.add("s.pdf", "pdf", text_pdf(STATEMENT))
    assert hk.pdf_tables(fid) == []  # plain text, no ruled table
    with pytest.raises(ValueError):
        hk.pdf_text(job.add("t.csv", "csv", b"a,b\n1,2\n"))


# ------------------------------------------------------------------ fix round 1


def test_chart_fields_escape_dots_and_brackets(job):
    slash = r"a\b"  # a column name with a real backslash in it
    df = pd.DataFrame({"price.gbp": [1.0, 2.0], "Amount [GBP]": [3.0, 4.0], slash: ["x", "y"]})
    hk.scatter(df, "price.gbp", "Amount [GBP]", color=slash, title="Escaped")
    hk.line(df, "price.gbp", "Amount [GBP]", title="Marked", highlight=[{"x": 1.5, "label": "m"}])
    spec, marked = (c["spec"] for c in by_type(job.results(), "chart"))
    main = spec["layer"][0]
    assert main["encoding"]["x"]["field"] == r"price\.gbp"
    assert main["encoding"]["y"]["field"] == r"Amount \[GBP\]"
    assert main["encoding"]["color"]["field"] == r"a\\b"
    assert [t["field"] for t in main["encoding"]["tooltip"]] == [
        r"price\.gbp",
        r"Amount \[GBP\]",
        r"a\\b",
    ]
    legend = next(p for p in main["params"] if p["name"] == "legend")
    assert legend["select"]["fields"] == [r"a\\b"]
    assert spec["data"]["values"][0] == {"price.gbp": 1.0, "Amount [GBP]": 3.0, slash: "x"}
    # Axis, legend and tooltip titles show the raw column names, not the escapes.
    assert main["encoding"]["x"]["title"] == "price.gbp"
    assert main["encoding"]["y"]["title"] == "Amount [GBP]"
    assert main["encoding"]["color"]["title"] == slash
    assert [t["title"] for t in main["encoding"]["tooltip"]] == [
        "price.gbp",
        "Amount [GBP]",
        slash,
    ]
    rule = marked["layer"][1]
    assert rule["encoding"]["x"]["field"] == r"price\.gbp"
    assert rule["encoding"]["x"]["title"] == "price.gbp"
    assert rule["data"]["values"][0] == {"price.gbp": 1.5, "hk_label": "m"}


def test_row_numbers_follow_filter_and_sort_only(job):
    data = "n,group\n" + "".join(f"{i},{'ab'[i % 2]}\n" for i in range(10))
    df = hk.read(job.add("rows.csv", "csv", data.encode()))
    assert df.index.name == "hk_row"
    hk.table("Filtered", df[df["n"] >= 7])
    hk.table("Sorted", df.sort_values("n", ascending=False).head(3))
    hk.table("Reset", df.reset_index(drop=True))
    hk.table("Grouped", df.groupby("group").agg(total=("n", "sum")))
    tables = {t["name"]: t for t in by_type(job.results(), "table")}
    assert tables["Filtered"]["row_numbers"] == [9, 10, 11]
    assert tables["Sorted"]["row_numbers"] == [11, 10, 9]
    assert tables["Reset"]["row_numbers"] == []
    assert tables["Grouped"]["row_numbers"] == []


def test_numpy_dates_and_durations_are_not_integers(job):
    hk.result("when", np.datetime64("2026-01-05T10:00:00.000000000"))
    hk.result("gap", np.timedelta64(90, "m"))
    hk.table("Dates", pd.DataFrame({"d": [np.datetime64("2026-01-05T00:00:00.000000000")]}))
    hk.table("Gaps", pd.DataFrame({"g": [np.timedelta64(90, "m"), pd.Timedelta(days=1)]}))
    values = {r["name"]: r for r in by_type(job.results(), "value")}
    assert values["when"]["value"].startswith("2026-01-05T10:00")
    assert values["when"]["is_date"] is True
    assert values["gap"]["value"] == "P0DT1H30M0S"
    assert values["gap"]["is_date"] is False
    dates, gaps = by_type(job.results(), "table")
    assert dates["rows"] == [["2026-01-05T00:00:00"]]
    assert gaps["rows"] == [["P0DT1H30M0S"], ["P1DT0H0M0S"]]  # same format as hk.result


def test_profile_knows_day_month_name_year_with_dashes():
    from hk._profile import _date_format

    assert _date_format("12-Mar-2026") == "D Month YYYY"
    assert _date_format("12 March, 2026") == "D Month YYYY"
    assert _date_format("12 Mar#2026") is None  # "#" is not a separator after the month


def test_read_semicolon_dot_thousands_integers(job):
    data = "Qty;Price;Code\n1.500;2,5;1.5\n20;3;2.25\n12.000.000;4,75;3\n"
    df = hk.read(job.add("q.csv", "csv", data.encode()))
    assert df["Qty"].tolist() == [1500, 20, 12_000_000]
    assert df["Price"].tolist() == [2.5, 3.0, 4.75]
    assert df["Code"].tolist() == [1.5, 2.25, 3.0]  # not all thousands-shaped: dots are decimals
    assert by_type(job.results(), "cleaning") == [
        {
            "type": "cleaning",
            "note": 'Read "1.500" style numbers in column Qty as thousands (1.500 = 1500).',
        }
    ]


def test_reading_notes_are_emitted_once_and_never_left_behind(job):
    from hk._read import read_text_table

    fid = job.add("q.csv", "csv", b"Qty;Name\n1.500;a\n20;b\n")
    path = job.manifest["files"][0]["path"]
    report = inspect(path, "csv")
    df, _ = read_text_table(path)
    assert report["ok"] is True
    assert job.results() == []  # inspecting or the low-level reader report nothing
    assert not any("note" in key for key in df.attrs)
    hk.read(fid)
    assert len(by_type(job.results(), "cleaning")) == 1
    assert not any("note" in key for key in hk.read(fid).attrs)


def test_read_semicolon_three_place_dot_decimals_stay_decimals(job):
    data = "Weight;Rate\n0.250;1.125\n1.125;0.500\n45.320;2\n"
    df = hk.read(job.add("w.csv", "csv", data.encode()))
    assert df["Weight"].tolist() == [0.25, 1.125, 45.32]
    assert df["Rate"].tolist() == [1.125, 0.5, 2.0]
    assert df["Weight"].dtype == float
    assert by_type(job.results(), "cleaning") == []


def test_read_semicolon_huge_dot_thousands_stay_float(job):
    data = "Big;Id\n12.000.000.000.000.000;1\n1.000;2\n"
    df = hk.read(job.add("big.csv", "csv", data.encode()))
    assert df["Big"].dtype == float  # over 2**53: never squeezed into int64
    assert df["Big"].tolist() == [1.2e16, 1000.0]


# ------------------------------------------------------------------ dates and gaps


def smoke_sales_csv() -> bytes:
    """The real-model smoke's file: 58 dated rows (North has no rows 12-24 March), 4 dates
    written DD/MM/YYYY among the YYYY-MM-DD ones, 3 negative amounts, 2 duplicates."""
    import random

    rng = random.Random(7)

    def amount() -> str:
        return f"{rng.randint(50, 500)}.{rng.randint(0, 99):02d}"

    rows: list[list[str]] = []
    for day in range(1, 32):
        for region in ("North", "South"):
            if region == "North" and 12 <= day <= 24:
                continue
            rows.append([f"2025-03-{day:02d}", region, amount()])
    for day in range(1, 10):
        rows.append([f"2025-03-{day:02d}", "East", amount()])
    rows.sort(key=lambda r: r[0])
    for index in (3, 17, 40):
        rows[index][2] = "-" + rows[index][2]
    for index in (8, 22, 35, 50):
        year, month, day = rows[index][0].split("-")
        rows[index][0] = f"{day}/{month}/{year}"
    for index in (5, 30):
        rows.append(list(rows[index]))
    lines = ["date,region,amount"] + [",".join(row) for row in rows]
    return ("\n".join(lines) + "\n").encode("utf-8")


NORTH_GAP_LINE = (
    "North has no rows from 2025-03-12 to 2025-03-24 (13 days), while other region values do."
)


def test_read_mixed_iso_and_day_first_dates_keeps_every_month(job):
    df = hk.read(job.add("sales.csv", "csv", smoke_sales_csv()))
    assert pd.api.types.is_datetime64_any_dtype(df["date"])
    assert df["date"].notna().all()
    assert set(df["date"].dt.month) == {3}  # never 2025-12-03 from a dayfirst ISO read
    written = df.loc[[8, 22, 35, 50], "date"].dt.strftime("%Y-%m-%d").tolist()
    assert written == ["2025-03-03", "2025-03-08", "2025-03-16", "2025-03-28"]  # day first
    notes = [r["note"] for r in by_type(job.results(), "cleaning")]
    assert notes == [
        "Read 4 dates in column date written DD/MM/YYYY (e.g. 03/03/2025) as day/month/year; "
        "the others are YYYY-MM-DD."
    ]


def test_gaps_finds_the_north_gap(job):
    df = hk.read(job.add("sales.csv", "csv", smoke_sales_csv()))
    found = hk.gaps(df, "date", by="region")
    # Each group runs from the file's first date to its last: East stops on 9 March.
    assert found.to_dict("records") == [
        {
            "group": "East",
            "last_before": dt.date(2025, 3, 9),
            "first_missing": dt.date(2025, 3, 10),
            "last_missing": dt.date(2025, 3, 31),
            "days": 22,
        },
        {
            "group": "North",
            "last_before": dt.date(2025, 3, 11),
            "first_missing": dt.date(2025, 3, 12),
            "last_missing": dt.date(2025, 3, 24),
            "days": 13,
        },
    ]
    records = job.results()
    tables = {t["name"]: t for t in by_type(records, "table")}
    gap_table = tables["Gaps in date by region"]
    assert gap_table["columns"] == ["group", "last_before", "first_missing", "last_missing", "days"]
    assert gap_table["rows"] == [
        ["East", "2025-03-09", "2025-03-10", "2025-03-31", 22],
        ["North", "2025-03-11", "2025-03-12", "2025-03-24", 13],
    ]
    assert "a gap at either end runs to that date" in gap_table["note"]
    notes = [r["note"] for r in by_type(records, "cleaning")]
    assert NORTH_GAP_LINE in notes
    assert (
        "East has no rows from 2025-03-10 to 2025-03-31 (22 days, to the end of the file), "
        "while other region values do." in notes
    )
    values = {r["name"]: r for r in by_type(records, "value")}
    assert values["North: first missing day"]["value"] == "2025-03-12"
    assert values["North: first missing day"]["is_date"] is True
    assert values["North: last missing day"]["value"] == "2025-03-24"
    assert values["North: missing days"]["value"] == 13
    assert values["Gaps found"]["value"] == 2
    assert values["East: missing days"]["value"] == 22


def test_gaps_reads_text_dates_and_caps_results(job):
    days = [f"2025-01-{d:02d}" for d in range(1, 32, 3)]  # 2 days missing between each
    df = pd.DataFrame({"when": days})
    found = hk.gaps(df, "when")
    assert len(found) == 10 and set(found["group"]) == {"all"}
    values = [r["name"] for r in by_type(job.results(), "value")]
    assert "all gap 10: missing days" in values and len(values) == 1 + 3 * 10
    assert hk.gaps(df, "when", min_days=3).empty


def test_dates_decides_day_or_month_first_from_the_column(job):
    month_first = hk.dates(pd.Series(["03/25/2025", "03/08/2025", "2025-03-01"], name="d"))
    days = month_first.dt.strftime("%Y-%m-%d").tolist()
    assert days == ["2025-03-25", "2025-03-08", "2025-03-01"]
    unsure = hk.dates(pd.Series(["08.03.2025", "09.03.2025"], name="d2"))
    assert unsure.dt.strftime("%Y-%m-%d").tolist() == ["2025-03-08", "2025-03-09"]
    both = hk.dates(pd.Series(["25/03/2025", "03/25/2025"], name="d3"))
    assert both.tolist() == ["25/03/2025", "03/25/2025"]
    bad = hk.dates(pd.Series(["31/02/2025", "28/02/2025", None], name="d4"))
    assert bad.isna().tolist() == [True, False, True]
    notes = [r["note"] for r in by_type(job.results(), "cleaning")]
    assert notes == [
        "Read 2 dates in column d written MM/DD/YYYY (e.g. 03/25/2025) as month/day/year; "
        "the others are YYYY-MM-DD.",
        "Read 2 dates in column d2 written DD.MM.YYYY (e.g. 08.03.2025) as day/month/year "
        "(no part is over 12, so day first was assumed).",
        "Column d3 has dates written day first (e.g. 25/03/2025) and month first "
        "(e.g. 03/25/2025), so it was left as text.",
        "Read 1 date in column d4 written DD/MM/YYYY (e.g. 28/02/2025) as day/month/year.",
        "1 value in column d4 could not be read as a date (e.g. 31/02/2025) and was left empty.",
    ]


def test_profile_still_lists_dates_written_another_way(job):
    hk.profile(job.add("sales.csv", "csv", smoke_sales_csv()))
    records = job.results()
    notes = " | ".join(r["note"] for r in by_type(records, "cleaning"))
    assert "mixes date formats: YYYY-MM-DD (56), DD/MM/YYYY or MM/DD/YYYY (4)" in notes
    odd = {t["name"]: t for t in by_type(records, "table")}["Dates not in the usual format in date"]
    assert odd["total_rows"] == 4


def test_result_with_the_value_first_says_the_name_comes_first(job):
    with pytest.raises(TypeError, match=r"the name comes first, e\.g\. hk\.result\('Total spent'"):
        hk.result(98.46, "amount")


def test_reading_the_same_file_twice_notes_once(job):
    fid = job.add("sales.csv", "csv", smoke_sales_csv())
    hk.read(fid)
    hk.read(fid)
    assert len(by_type(job.results(), "cleaning")) == 1


def test_read_leaves_codes_versions_and_dates_with_words_as_text(job):
    data = (
        "code,version,when,day\n"
        "10-20-3000,1.2.10,12/03/2025 x,2025-03-01\n"
        "12-34-5678,3.4.15,13/03/2025,2025-03-02\n"
        "10-20-3000,2.0.11,14/03/2025,2025-03-03\n"
    )
    df = hk.read(job.add("codes.csv", "csv", data.encode()))
    assert df["code"].tolist() == ["10-20-3000", "12-34-5678", "10-20-3000"]
    assert df["version"].tolist() == ["1.2.10", "3.4.15", "2.0.11"]
    assert df["when"].tolist() == ["12/03/2025 x", "13/03/2025", "14/03/2025"]
    assert pd.api.types.is_datetime64_any_dtype(df["day"])
    assert by_type(job.results(), "cleaning") == []  # nothing was read as a date


def test_read_leaves_dates_that_do_not_all_parse_as_text(job):
    data = "a,b\n31/02/2025,1.2.2010\n28/02/2025,3.4.2011\n"
    df = hk.read(job.add("odd.csv", "csv", data.encode()))
    assert df["a"].tolist() == ["31/02/2025", "28/02/2025"]  # 31 February: kept as written
    assert pd.api.types.is_datetime64_any_dtype(df["b"])  # whole dates with a 4-digit year


def test_dates_note_never_says_yy_for_a_three_digit_year(job):
    out = hk.dates(pd.Series(["01/02/025", "13/02/2025"], name="d"))
    assert out.isna().tolist() == [True, False]
    notes = [r["note"] for r in by_type(job.results(), "cleaning")]
    assert notes == [
        "Read 1 date in column d written DD/MM/YYYY (e.g. 13/02/2025) as day/month/year.",
        "1 value in column d could not be read as a date (e.g. 01/02/025) and was left empty.",
    ]


def test_profile_finds_the_north_gap_by_itself(job):
    hk.profile(job.add("sales.csv", "csv", smoke_sales_csv()))
    records = job.results()
    notes = [r["note"] for r in by_type(records, "cleaning")]
    assert NORTH_GAP_LINE in notes
    values = {r["name"]: r["value"] for r in by_type(records, "value")}
    assert values["North: first missing day"] == "2025-03-12"
    assert values["North: last missing day"] == "2025-03-24"
    assert values["North: missing days"] == 13
    tables = {t["name"] for t in by_type(records, "table")}
    assert "Gaps in date by region" in tables


def _region_csv(stops: dict[str, int]) -> bytes:
    lines = ["date,region,amount"]
    for day in range(1, 32):
        for region in ("North", "East"):
            if day <= stops.get(region, 31):
                lines.append(f"2025-03-{day:02d},{region},{day}.50")
    return ("\n".join(lines) + "\n").encode()


def test_profile_reports_a_group_that_stops_early(job):
    hk.profile(job.add("east.csv", "csv", _region_csv({"East": 19})))
    records = job.results()
    notes = [r["note"] for r in by_type(records, "cleaning")]
    assert (
        "East has no rows from 2025-03-20 to 2025-03-31 (12 days, to the end of the file), "
        "while other region values do." in notes
    )
    values = {r["name"]: r["value"] for r in by_type(records, "value")}
    assert values["East: first missing day"] == "2025-03-20"
    assert values["East: missing days"] == 12


def test_profile_reports_no_gaps_when_there_are_none(job):
    hk.profile(job.add("full.csv", "csv", _region_csv({})))
    records = job.results()
    assert not any("has no rows" in r["note"] for r in by_type(records, "cleaning"))
    assert not any(t["name"].startswith("Gaps in") for t in by_type(records, "table"))
    assert not any("missing day" in r["name"] for r in by_type(records, "value"))


def _bank_statement_csv() -> bytes:
    """150 card and bank lines over a year: a few days a week, five transaction types."""
    types = ["DEB", "CR", "DD", "SO", "FPI"]
    start = dt.date(2025, 1, 1)
    lines = ["Date,Type,Amount"]
    for n in range(150):
        day = start + dt.timedelta(days=(n * 12) // 5)  # rows on about 2 days in 5
        lines.append(f"{day.isoformat()},{types[(n * 7) % 5]},{(n % 40) - 20}.15")
    return ("\n".join(lines) + "\n").encode()


def test_profile_reports_no_gaps_in_a_sparse_bank_statement(job):
    hk.profile(job.add("statement.csv", "csv", _bank_statement_csv()))
    records = job.results()
    assert not any("has no rows" in r["note"] or "more gap" in r["note"]
                   for r in by_type(records, "cleaning"))  # fmt: skip
    assert not any("missing day" in r["name"] for r in by_type(records, "value"))
    assert not any(t["name"].startswith("Gaps in") for t in by_type(records, "table"))


def test_profile_still_reports_north_and_east_in_daily_sales(job):
    hk.profile(job.add("sales.csv", "csv", smoke_sales_csv()))
    notes = [r["note"] for r in by_type(job.results(), "cleaning")]
    assert NORTH_GAP_LINE in notes
    assert (
        "East has no rows from 2025-03-10 to 2025-03-31 (22 days, to the end of the file), "
        "while other region values do." in notes
    )


def test_profile_caps_gap_lines_at_ten(job):
    # 12 shops, each closed for one 3-day spell at a different time; every day has rows.
    lines = ["date,shop,amount"]
    for day in range(1, 61):
        for shop in range(12):
            closed = 4 * shop + 3 <= day <= 4 * shop + 5
            if not closed:
                when = dt.date(2025, 4, 1) + dt.timedelta(days=day - 1)
                lines.append(f"{when.isoformat()},Shop {shop:02d},1.00")
    hk.profile(job.add("shops.csv", "csv", ("\n".join(lines) + "\n").encode()))
    records = job.results()
    gap_lines = [r["note"] for r in by_type(records, "cleaning") if "has no rows" in r["note"]]
    assert len(gap_lines) == 10
    assert "…and 2 more gaps (see the table)." in [r["note"] for r in by_type(records, "cleaning")]
    [gap_table] = [t for t in by_type(records, "table") if t["name"] == "Gaps in date by shop"]
    assert gap_table["total_rows"] == 12
    assert len([r for r in by_type(records, "value") if r["name"].endswith("missing days")]) == 10


def test_result_takes_a_dict_one_result_per_item(job):
    hk.result("Null counts", {"date": 0, "region": 2, "amount": np.int64(1)})
    values = [(r["name"], r["value"]) for r in by_type(job.results(), "value")]
    assert values == [("Null counts: date", 0), ("Null counts: region", 2),
                      ("Null counts: amount", 1)]  # fmt: skip
    assert by_type(job.results(), "cleaning") == []


def test_result_and_document_take_a_series(job):
    df = pd.DataFrame({"a": [1.0, None], "b": [None, None]})
    hk.result("Missing", df.isnull().sum(), unit="rows")
    hk.document("Totals", pd.Series({"Opening": "1,200.00", "Closing": "980.10"}), page=2)
    records = by_type(job.results(), "value")
    assert [(r["name"], r["value"], r["unit"]) for r in records[:2]] == [
        ("Missing: a", 1, "rows"), ("Missing: b", 2, "rows"),
    ]  # fmt: skip
    assert [(r["name"], r["kind"], r["page"]) for r in records[2:]] == [
        ("Totals: Opening", "document", 2), ("Totals: Closing", "document", 2),
    ]  # fmt: skip


def test_result_caps_a_big_dict_at_20_with_a_note(job):
    hk.result("Per day", {f"d{i:02d}": i for i in range(25)})
    names = [r["name"] for r in by_type(job.results(), "value")]
    assert names == [f"Per day: d{i:02d}" for i in range(20)]
    assert [r["note"] for r in by_type(job.results(), "cleaning")] == [
        "Reported the first 20 of 25 values for Per day; use hk.table for all of them."
    ]


def test_result_skips_values_that_are_not_single_values(job):
    hk.result("Mixed", {"ok": 3, "nested": {"x": 1}, "listed": [1, 2]})
    assert [r["name"] for r in by_type(job.results(), "value")] == ["Mixed: ok"]
    assert [r["note"] for r in by_type(job.results(), "cleaning")] == [
        "Skipped 2 values for Mixed that aren't single values (nested, listed); "
        "use hk.table for them."
    ]


def test_result_still_refuses_a_table_or_a_list(job):
    for value in (pd.DataFrame({"a": [1]}), [1, 2]):
        with pytest.raises(TypeError, match="use hk.table for many values"):
            hk.result("Many", value)


def test_long_result_names_stay_within_100_characters(job):
    key = "Payment to " + "x" * 139  # 150 characters
    hk.result("Spending by description", {key: 12.5, "Rent": 900})
    hk.result("n" * 150, 1)
    [long, rent, plain] = by_type(job.results(), "value")
    assert len(long["name"]) == 100 and long["name"].endswith("…")
    assert long["name"].startswith("Spending by description: Payment to xxx")
    assert long["value"] == 12.5
    assert rent["name"] == "Spending by description: Rent"
    assert len(plain["name"]) == 100 and plain["name"].endswith("…")


def test_an_empty_dict_or_series_says_so(job):
    hk.result("Refunds", {})
    hk.result("Fees", pd.Series(dtype=float))
    assert by_type(job.results(), "value") == []
    assert [r["note"] for r in by_type(job.results(), "cleaning")] == [
        "Refunds: no values",
        "Fees: no values",
    ]
