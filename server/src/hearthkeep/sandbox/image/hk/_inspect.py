"""What is inside one file: sheets, rows, columns, pages, and whether it needs a password."""

from __future__ import annotations

from pathlib import Path

from hk._pdf import PasswordProtected, is_password_error, open_pdf
from hk._read import (
    _filled,
    _width,
    column_names,
    find_header,
    read_json,
    read_text_table,
)


def _blank_report(kind: str) -> dict:
    return {
        "ok": True,
        "problem": None,
        "kind": kind,
        "pages": None,
        "scanned_pages": [],
        "sheets": [],
        "rows": None,
        "columns": [],
        "encoding": None,
        "delimiter": None,
    }


def _text(report: dict, path: Path) -> None:
    df, table = read_text_table(path)
    report["encoding"] = table.encoding
    report["delimiter"] = table.delimiter
    report["rows"] = int(len(df))
    report["columns"] = [str(c) for c in df.columns]


def _first_rows(rows_iter, limit: int = 26) -> list[list]:
    rows = []
    for row in rows_iter:
        rows.append(list(row))
        if len(rows) >= limit:
            break
    return rows


def _header_of(rows: list[list]) -> tuple[int, list[str]]:
    if not any(_filled(r) for r in rows):
        return 0, []
    h = find_header(rows, padded=True)
    width = max(_width(r) for r in rows)
    return h, column_names((list(rows[h]) + [None] * width)[:width])


def _xlsx(report: dict, path: Path) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for i, ws in enumerate(wb.worksheets):
            if ws.max_row is None or ws.max_column is None:
                ws.calculate_dimension(force=True)
            rows, cols = int(ws.max_row or 0), int(ws.max_column or 0)
            report["sheets"].append({"name": ws.title, "rows": rows, "columns": cols})
            if i == 0:
                first = _first_rows(ws.iter_rows(values_only=True))
                h, names = _header_of(first)
                report["columns"] = names
                report["rows"] = max(rows - (h + 1), 0) if names else 0
    finally:
        wb.close()


def _xls(report: dict, path: Path) -> None:
    import xlrd

    book = xlrd.open_workbook(str(path), on_demand=True)
    try:
        for i, name in enumerate(book.sheet_names()):
            ws = book.sheet_by_index(i)
            report["sheets"].append({"name": name, "rows": ws.nrows, "columns": ws.ncols})
            if i == 0:
                first = [ws.row_values(r) for r in range(min(ws.nrows, 26))]
                h, names = _header_of(first)
                report["columns"] = names
                report["rows"] = max(ws.nrows - (h + 1), 0) if names else 0
            book.unload_sheet(i)
    finally:
        book.release_resources()


def _pdf(report: dict, path: Path) -> None:
    with open_pdf(path) as pdf:
        report["pages"] = len(pdf.pages)
        for n, page in enumerate(pdf.pages, start=1):
            if not (page.extract_text() or "").strip():
                report["scanned_pages"].append(n)
            page.close()


def inspect(path: str | Path, kind: str) -> dict:
    """The `inspect.json` report for one file. Never raises: a file that can't be read is
    `ok: false` with `problem` "password" or "unreadable"."""
    report = _blank_report(kind)
    path = Path(path)
    try:
        if kind in ("csv", "txt"):
            _text(report, path)
        elif kind == "xlsx":
            _xlsx(report, path)
        elif kind == "xls":
            _xls(report, path)
        elif kind == "json":
            df = read_json(path)
            report["rows"] = int(len(df))
            report["columns"] = [str(c) for c in df.columns]
        elif kind == "pdf":
            _pdf(report, path)
        else:
            raise ValueError(f"unknown kind {kind!r}")
    except PasswordProtected:
        report.update(ok=False, problem="password")
    except Exception as exc:  # anything else: the file can't be read
        problem = "password" if kind == "xls" and is_password_error(exc) else "unreadable"
        report.update(ok=False, problem=problem)
    return report
