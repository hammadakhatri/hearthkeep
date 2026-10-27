"""Reading the chat's files into pandas, the way a careful person would open them."""

from __future__ import annotations

import csv
import datetime as _dt
import io
import json
import os
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd

from hk._dates import FORMATS_ATTR, auto_dates
from hk._results import ROW_INDEX, note_once

DELIMITERS = (",", ";", "\t", "|")
HEADER_SCAN_ROWS = 20
DECIMAL_COMMA = re.compile(r"[-+]?\d+(\.\d{3})*,\d+")
DOT_THOUSANDS = re.compile(r"[-+]?[1-9]\d{0,2}(\.\d{3})+")
INT64_SAFE = 2**53  # larger whole numbers stay float rather than risk int64
PLAIN_INTEGER = re.compile(r"[-+]?\d+")
_SAMPLE_CHARS = 64 * 1024


def job_dir() -> Path:
    return Path(os.environ.get("HK_JOB", "/job"))


def files() -> list[dict]:
    """The chat's files: `[{"id", "name", "kind", "path"}]`."""
    manifest = json.loads((job_dir() / "manifest.json").read_text(encoding="utf-8"))
    return [dict(f) for f in manifest.get("files", [])]


def find(file: str) -> dict:
    """A file by its id, else by its name."""
    entries = files()
    for key in ("id", "name"):
        for entry in entries:
            if entry.get(key) == file:
                return entry
    known = ", ".join(f"{e.get('id')} ({e.get('name')})" for e in entries) or "none"
    raise ValueError(f"no file {file!r}; the chat's files are: {known}")


# ---------------------------------------------------------------- text files


@dataclass
class TextTable:
    """A delimited text file parsed into records, with what was learned on the way."""

    encoding: str
    delimiter: str | None
    records: list[list[str]] = field(default_factory=list)
    lines: list[int] = field(default_factory=list)  # 1-based file line of each record
    notes: list[str] = field(default_factory=list)  # what typing the columns assumed


def decode(data: bytes) -> tuple[str, str]:
    """(text, encoding): UTF-8 (BOM or not), UTF-16 (BOM or the NUL pattern), else cp1252."""
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace"), "utf-8-sig"
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace").lstrip("\ufeff"), "utf-16"
    sample = data[:4096]
    if b"\x00" in sample:
        odd = sample[1::2]
        enc = "utf-16-le" if odd.count(0) > len(odd) // 2 else "utf-16-be"
        return data.decode(enc, errors="replace"), "utf-16"
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace"), "cp1252"


def sniff_delimiter(text: str) -> str | None:
    """The delimiter whose field count is the most consistent across the first lines."""
    lines = [ln for ln in text[:_SAMPLE_CHARS].splitlines() if ln.strip()][:40]
    best: tuple[int, int, int] | None = None
    choice = None
    for order, delim in enumerate(DELIMITERS):
        counts = [len(r) for r in csv.reader(lines, delimiter=delim)]
        wide = [c for c in counts if c > 1]
        if not wide:
            continue
        mode, agree = Counter(wide).most_common(1)[0]
        score = (agree, mode, -order)
        if best is None or score > best:
            best, choice = score, delim
    return choice


def parse_text(data: bytes) -> TextTable:
    text, encoding = decode(data)
    delimiter = sniff_delimiter(text)
    table = TextTable(encoding=encoding, delimiter=delimiter)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter or "\x1f")
    start = 1
    for record in reader:
        if any(cell.strip() for cell in record):
            table.records.append(record)
            table.lines.append(start)
        start = reader.line_num + 1
    return table


# ---------------------------------------------------------------- header detection


def _filled(row) -> int:
    return sum(1 for v in row if not _blank(v))


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _width(row) -> int:
    width = 0
    for i, v in enumerate(row):
        if not _blank(v):
            width = i + 1
    return width


def find_header(rows: list[list], *, padded: bool) -> int:
    """Index of the header among the first 20 rows: the first row with at least two filled
    cells whose following rows keep its column count. `padded` rows (spreadsheets) are
    compared by the last filled cell, delimited text by the number of fields."""
    size = _width if padded else len
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        if _filled(row) < 2:
            continue
        following = [size(r) for r in rows[i + 1 : i + 6]]
        if not following:
            return i
        if max(following) <= size(row) and size(row) in following:
            return i
    return 0


def column_names(header: list) -> list[str]:
    names: list[str] = []
    seen: Counter[str] = Counter()
    for i, raw in enumerate(header):
        name = "" if raw is None else str(raw).strip().lstrip("\ufeff")
        if not name:
            name = f"column_{i + 1}"
        seen[name] += 1
        if seen[name] > 1:
            name = f"{name}_{seen[name]}"
        names.append(name)
    return names


# ---------------------------------------------------------------- typing columns


def _text_column(
    values: pd.Series, delimiter: str | None, notes: list[str] | None = None
) -> pd.Series:
    """Type one text column. Dot-thousands conversions are named in `notes`."""
    stripped = values.map(lambda v: v.strip() if isinstance(v, str) else v)
    stripped = stripped.mask(stripped.map(_blank), np.nan)
    present = stripped.dropna()
    if present.empty:
        return stripped
    texts = present.astype(str)
    if delimiter == ";":
        # European numbers: "1.234,5" (decimal comma) and "1.500" (dot thousands), when
        # every value in the column has one of those shapes or is a plain integer.
        comma = texts.map(lambda s: bool(DECIMAL_COMMA.fullmatch(s)))
        dots = texts.map(lambda s: bool(DOT_THOUSANDS.fullmatch(s)))
        plain_int = texts.map(lambda s: bool(PLAIN_INTEGER.fullmatch(s)))
        if (comma | dots).any() and (comma | dots | plain_int).all():
            converted = stripped.map(
                lambda s: float(s.replace(".", "").replace(",", ".")) if isinstance(s, str) else s
            ).astype(float)
            if dots.any() and notes is not None:
                notes.append(
                    f'Read "1.500" style numbers in column {values.name} as thousands '
                    "(1.500 = 1500)."
                )
            whole = converted.notna().all() and (converted.abs() <= INT64_SAFE).all()
            if not comma.any() and whole:
                return converted.astype("int64")
            return converted
    numbers = pd.to_numeric(stripped, errors="coerce")
    if numbers.notna().sum() == len(present):
        return numbers
    # Whole dates (YYYY-MM-DD, DD/MM/YYYY and the like, even mixed; never dayfirst on ISO
    # dates), only when every value parses: anything else stays text exactly as it was.
    dated = auto_dates(stripped.rename(values.name), notes)
    return stripped if dated is None else dated


def _cell_column(values: pd.Series) -> pd.Series:
    present = values.dropna()
    present = present[~present.map(_blank)]
    if present.empty:
        return values.mask(values.map(_blank), np.nan)
    filled = values.mask(values.map(_blank), np.nan)
    kinds = {type(v) for v in present}
    if kinds <= {int, float}:
        return pd.to_numeric(filled)
    if all(isinstance(v, _dt.date) for v in present):
        try:
            return pd.to_datetime(filled)
        except (TypeError, ValueError, OverflowError):
            return filled
    return filled


def _frame(names: list[str], body: list[list], lines: list[int]) -> pd.DataFrame:
    width = len(names)
    fixed = [(list(r) + [None] * width)[:width] for r in body]
    df = pd.DataFrame(fixed, columns=names, dtype=object)
    first = lines[0] if lines else None
    if lines and lines[-1] - first != len(lines) - 1:
        # Blank lines were skipped: keep index + first data row = the file row.
        df.index = pd.Index([ln - first for ln in lines], dtype="int64")
    df.attrs["hk_first_data_row"] = first
    return df


# ---------------------------------------------------------------- per kind


def read_text_table(path: Path, *, header: int | None = None) -> tuple[pd.DataFrame, TextTable]:
    table = parse_text(Path(path).read_bytes())
    records = table.records
    if not records:
        df = pd.DataFrame()
        df.attrs["hk_first_data_row"] = None
        return df, table
    h = find_header(records, padded=False) if header is None else header
    names = column_names(records[h])
    df = _frame(names, records[h + 1 :], table.lines[h + 1 :])
    if not table.lines[h + 1 :]:
        df.attrs["hk_first_data_row"] = table.lines[h] + 1
    for col in df.columns:
        typed = _text_column(df[col], table.delimiter, table.notes)
        formats = typed.attrs.pop(FORMATS_ATTR, None)
        df[col] = typed
        if formats:  # hk.profile lists the rows written another way
            df.attrs.setdefault(FORMATS_ATTR, {})[col] = formats
    df.attrs["hk_header_row"] = table.lines[h]
    df.attrs["hk_encoding"] = table.encoding
    df.attrs["hk_delimiter"] = table.delimiter
    return df, table


def _sheet_frame(rows: list[list], row_numbers: list[int], header: int | None) -> pd.DataFrame:
    while rows and _filled(rows[-1]) == 0:
        rows.pop()
        row_numbers.pop()
    if not rows:
        df = pd.DataFrame()
        df.attrs["hk_first_data_row"] = None
        return df
    h = find_header(rows, padded=True) if header is None else header
    width = max(_width(rows[h]), max((_width(r) for r in rows[h + 1 :]), default=0))
    names = column_names((list(rows[h]) + [None] * width)[:width])
    body = rows[h + 1 :]
    df = _frame(names, body, row_numbers[h + 1 :])
    if not body:
        df.attrs["hk_first_data_row"] = row_numbers[h] + 1
    for col in df.columns:
        df[col] = _cell_column(df[col])
    df.attrs["hk_header_row"] = row_numbers[h]
    return df


def read_xlsx(path: Path, *, sheet=None, header: int | None = None) -> pd.DataFrame:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = _pick_sheet(wb.worksheets, wb.sheetnames, sheet)
        rows: list[list] = []
        numbers: list[int] = []
        for n, row in enumerate(ws.iter_rows(values_only=True), start=1):
            rows.append(list(row))
            numbers.append(n)
        return _sheet_frame(rows, numbers, header)
    finally:
        wb.close()


def read_xls(path: Path, *, sheet=None, header: int | None = None) -> pd.DataFrame:
    import xlrd

    book = xlrd.open_workbook(str(path), on_demand=True)
    try:
        names = book.sheet_names()
        if sheet is None:
            ws = book.sheet_by_index(0)
        elif isinstance(sheet, int):
            ws = book.sheet_by_index(sheet)
        elif sheet in names:
            ws = book.sheet_by_name(sheet)
        else:
            raise ValueError(f"no sheet {sheet!r}; the sheets are: {', '.join(names)}")
        rows = []
        for r in range(ws.nrows):
            row = []
            for c in range(ws.ncols):
                cell = ws.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    row.append(bool(cell.value))
                else:
                    row.append(cell.value)
            rows.append(row)
        return _sheet_frame(rows, list(range(1, len(rows) + 1)), header)
    finally:
        book.release_resources()


def _pick_sheet(sheets, names, sheet):
    if sheet is None:
        return sheets[0]
    if isinstance(sheet, int):
        return sheets[sheet]
    if sheet in names:
        return sheets[names.index(sheet)]
    raise ValueError(f"no sheet {sheet!r}; the sheets are: {', '.join(names)}")


def read_json(path: Path) -> pd.DataFrame:
    data = json.loads(decode(Path(path).read_bytes())[0])
    if isinstance(data, dict):
        lists = [k for k, v in data.items() if isinstance(v, list)]
        if len(lists) == 1 and all(isinstance(x, dict) for x in data[lists[0]]):
            data = data[lists[0]]
    df = pd.json_normalize(data)
    df.attrs["hk_first_data_row"] = None
    return df


def read(file: str, *, sheet: str | int | None = None, header: int | None = None) -> pd.DataFrame:
    """One of the chat's files (by id or name) as a DataFrame.

    CSV/text: encoding, delimiter, decimal commas and the header row are found for you.
    Excel: stored values (never formula text). JSON: flattened with json_normalize.
    `header` is the 0-based row of the header among the file's non-blank rows.
    df.attrs["hk_first_data_row"] is the file row of the first data row.
    """
    entry = find(file)
    kind, path = entry.get("kind"), Path(entry["path"])
    notes: list[str] = []
    if kind in ("csv", "txt"):
        df, table = read_text_table(path, header=header)
        notes = table.notes
    elif kind == "xlsx":
        df = read_xlsx(path, sheet=sheet, header=header)
    elif kind == "xls":
        df = read_xls(path, sheet=sheet, header=header)
    elif kind == "json":
        df = read_json(path)
    elif kind == "pdf":
        raise ValueError("a PDF is read with hk.pdf_text or hk.pdf_tables")
    else:
        raise ValueError(f"hk.read can't read a {kind} file")
    df.index.name = ROW_INDEX  # hk.table keeps file row numbers while this index survives
    for note in notes:
        note_once(note)  # what reading assumed is part of "How I got this", said once
    return df
